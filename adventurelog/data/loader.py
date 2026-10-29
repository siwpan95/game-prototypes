# -*- coding: utf-8 -*-
"""数据加载器 — 优先 Excel，回退 JSON"""

import json
import os
from typing import Dict, List, Optional, Any

from core.models import (JobDef, Skill, Equipment, Monster, Dungeon,
                          Race, Consumable, ExploreEvent)


class DataLoader:
    """从 config/game_data.xlsx 或 data/*.json 加载游戏数据"""

    def __init__(self):
        self.jobs: Dict[str, JobDef] = {}
        self.skills: Dict[str, Skill] = {}
        self.equipment: Dict[str, Equipment] = {}
        self.monsters: Dict[str, Monster] = {}
        self.dungeons: List[Dungeon] = []
        self.races: Dict[str, Race] = {}
        self.consumables: Dict[str, Consumable] = {}
        self.events: Dict[str, ExploreEvent] = {}
        self.narratives: List[str] = []

    def load_all(self, data_dir: str, config_dir: str) -> bool:
        """优先尝试从 Excel 加载，失败则从 JSON 加载"""
        xlsx_path = os.path.join(config_dir, "game_data.xlsx")
        if os.path.exists(xlsx_path):
            try:
                self._from_excel(xlsx_path)
                print("[Loader] 从 Excel 加载完成")
                return True
            except Exception as e:
                print(f"[Loader] Excel 加载失败 ({e})，回退 JSON...")

        self._from_json(data_dir)
        print("[Loader] 从 JSON 加载完成")
        return False

    # ---- JSON 加载 ----
    def _from_json(self, data_dir: str):
        self.jobs = self._load_json_dict(data_dir, "jobs.json", self._parse_job)
        self.skills = self._load_json_dict(data_dir, "skills.json", self._parse_skill)
        self.equipment = self._load_json_dict(data_dir, "equipment.json", self._parse_equip)
        self.monsters = self._load_json_dict(data_dir, "monsters.json", self._parse_monster)
        self.dungeons = self._load_json_list(data_dir, "dungeons.json", self._parse_dungeon)
        self.races = self._load_json_dict(data_dir, "races.json", self._parse_race)
        self.consumables = self._load_json_dict(data_dir, "consumes.json", self._parse_consumable)
        self.events = self._load_json_dict(data_dir, "events.json", self._parse_event)
        self.narratives = self._load_narratives(data_dir)

    def _from_json_partial(self, data_dir: str, which: list):
        """从 JSON 补充加载缺失的数据类型"""
        for key in which:
            if key == "races":
                self.races = self._load_json_dict(data_dir, "races.json", self._parse_race)
            elif key == "consumables":
                self.consumables = self._load_json_dict(data_dir, "consumes.json", self._parse_consumable)
            elif key == "events":
                self.events = self._load_json_dict(data_dir, "events.json", self._parse_event)
            elif key == "narratives":
                self.narratives = self._load_narratives(data_dir)

    def _load_json_dict(self, data_dir, filename, parser) -> dict:
        path = os.path.join(data_dir, filename)
        if not os.path.exists(path):
            return {}
        with open(path, "r", encoding="utf-8") as f:
            items = json.load(f)
        return {k: parser(k, v) for k, v in items.items()}

    def _load_json_list(self, data_dir, filename, parser) -> list:
        path = os.path.join(data_dir, filename)
        if not os.path.exists(path):
            return []
        with open(path, "r", encoding="utf-8") as f:
            items = json.load(f)
        return [parser(i) for i in items]

    @staticmethod
    def _parse_job(key, d):
        return JobDef(d["name"], d["emoji"], d["color"],
                      d["base_hp"], d["base_atk"], d["base_def"], d["base_spd"],
                      d.get("initial_skills", []))

    @staticmethod
    def _parse_skill(key, d):
        _jl = d.get("job_limit") or ""
        if isinstance(_jl, (list, tuple)):
            _jl = ",".join(str(s) for s in _jl)
        job_limit = [s.strip() for s in _jl.split(",") if s.strip()]
        stat_req_raw = d.get("stat_req", {}) or {}
        stat_req = {k: int(v) for k, v in stat_req_raw.items()} if isinstance(stat_req_raw, dict) else {}
        return Skill(d["name"], d["desc"], d["skill_type"], d["power"],
                     d.get("mp_cost", 0), d.get("cooldown", 0),
                     job_limit=job_limit,
                     level_req=int(d.get("level_req", 0)),
                     stat_req=stat_req)

    @staticmethod
    def _parse_equip(key, d):
        return Equipment(d["name"], d["slot"],
                         d.get("hp_bonus", 0), d.get("atk_bonus", 0),
                         d.get("def_bonus", 0), d.get("spd_bonus", 0),
                         d.get("tier", 1))

    @staticmethod
    def _parse_monster(key, d):
        return Monster(d["name"], d["hp"], d["atk"], d["def_"], d["spd"],
                       d["gold"], d["exp"],
                       d.get("emoji", "👾"), d.get("color", "#ff4444"),
                       d.get("skills", []))

    @staticmethod
    def _parse_dungeon(d):
        return Dungeon(d["name"], d["floors"],
                       d["monster_names"], d["boss_name"], d["difficulty"],
                       d.get("gold_mult", 1.0), d.get("exp_mult", 1.0),
                       d.get("event_chance", 0.25), d.get("trap_chance", 0.15))

    @staticmethod
    def _parse_race(key, d):
        return Race(d["name"], d["desc"],
                    d.get("hp_mod", 0), d.get("atk_mod", 0),
                    d.get("def_mod", 0), d.get("spd_mod", 0),
                    d.get("mp_mod", 0))

    @staticmethod
    def _parse_consumable(key, d):
        return Consumable(d["name"], d["desc"], d["item_type"],
                          d.get("effect_value", 0), d.get("price", 10),
                          d.get("emoji", "🧪"))

    @staticmethod
    def _parse_event(key, d):
        return ExploreEvent(d["name"], d["desc"], d["event_type"],
                            d.get("choices", []))

    @staticmethod
    def _load_narratives(data_dir: str) -> list:
        """加载叙事文字列表"""
        path = os.path.join(data_dir, "narratives.json")
        if not os.path.exists(path):
            return []
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data.get("narratives", [])

    # ---- Excel 加载 ----
    def _from_excel(self, xlsx_path: str):
        import openpyxl
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)

        self.jobs = self._read_sheet_dict(wb, "职业", self._row_to_job)
        self.skills = self._read_sheet_dict(wb, "技能", self._row_to_skill)
        self.equipment = self._read_sheet_dict(wb, "装备", self._row_to_equip)
        self.monsters = self._read_sheet_dict(wb, "怪物", self._row_to_monster)
        self.dungeons = self._read_sheet_list(wb, "地下城", self._row_to_dungeon)
        self.races = self._read_sheet_dict(wb, "种族", self._row_to_race)
        self.consumables = self._read_sheet_dict(wb, "消耗品", self._row_to_consumable)
        self.events = self._read_sheet_dict(wb, "事件", self._row_to_event)
        self.narratives = self._read_narratives(wb)

        # 回退: Excel 缺少新 sheet 时从 JSON 补充
        missing = []
        if not self.races:
            missing.append("races")
        if not self.consumables:
            missing.append("consumables")
        if not self.events:
            missing.append("events")
        if not self.narratives:
            missing.append("narratives")

        if missing:
            data_dir = os.path.join(os.path.dirname(xlsx_path), "..", "data")
            data_dir = os.path.normpath(data_dir)
            # fallback: CONFIG_DIR/../data  = project/data
            if not os.path.isdir(data_dir):
                # 从 xlsx_path 所在目录推导：config/../data/
                base = os.path.dirname(os.path.dirname(xlsx_path))
                data_dir = os.path.join(base, "data")
            self._from_json_partial(data_dir, missing)
            print(f"[Loader] JSON 补充加载: {', '.join(missing)}")

    def _read_sheet_dict(self, wb, sheet_name, row_parser) -> dict:
        if sheet_name not in wb.sheetnames:
            return {}
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        result = {}
        for row in rows:
            if row and row[0] is not None:
                key = str(row[0])
                result[key] = row_parser(row)
        return result

    def _read_sheet_list(self, wb, sheet_name, row_parser) -> list:
        if sheet_name not in wb.sheetnames:
            return []
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        return [row_parser(row) for row in rows if row and row[1] is not None]

    @staticmethod
    def _row_to_job(r):
        skills_str = str(r[7] or "") if len(r) > 7 else ""
        init_skills = [s.strip() for s in skills_str.split(",") if s.strip()]
        return JobDef(str(r[0]), str(r[1]), str(r[2]),
                      int(r[3]), int(r[4]), int(r[5]), int(r[6]), init_skills)

    @staticmethod
    def _row_to_skill(r):
        job_limit_str = str(r[6] or "") if len(r) > 6 else ""
        job_limit = [s.strip() for s in job_limit_str.split(",") if s.strip()]
        level_req = int(r[7] or 0) if len(r) > 7 else 0
        import json as _json
        stat_req = {}
        if len(r) > 8 and r[8]:
            try:
                stat_req = {k: int(v) for k, v in _json.loads(str(r[8])).items()}
            except Exception:
                stat_req = {}
        return Skill(str(r[0]), str(r[1]), str(r[2]), float(r[3]),
                     int(r[4] or 0), int(r[5] or 0),
                     job_limit=job_limit, level_req=level_req, stat_req=stat_req)

    @staticmethod
    def _row_to_equip(r):
        return Equipment(str(r[0]), str(r[1]),
                         int(r[2] or 0), int(r[3] or 0),
                         int(r[4] or 0), int(r[5] or 0),
                         int(r[6] or 1) if len(r) > 6 else 1)

    @staticmethod
    def _row_to_monster(r):
        skills_str = str(r[9] or "") if len(r) > 9 else ""
        monster_skills = [s.strip() for s in skills_str.split(",") if s.strip()]
        return Monster(str(r[0]), int(r[1]), int(r[2]), int(r[3]), int(r[4]),
                       int(r[5]), int(r[6]),
                       str(r[7] or "👾"), str(r[8] or "#ff4444"), monster_skills)

    @staticmethod
    def _row_to_dungeon(r):
        return Dungeon(str(r[0]), int(r[1]),
                       str(r[2]).split(","), str(r[3]), int(r[4]),
                       float(r[5] or 1.0), float(r[6] or 1.0),
                       float(r[7] or 0.25) if len(r) > 7 else 0.25,
                       float(r[8] or 0.15) if len(r) > 8 else 0.15)

    @staticmethod
    def _row_to_race(r):
        return Race(str(r[0]), str(r[1]),
                    int(r[2] or 0), int(r[3] or 0),
                    int(r[4] or 0), int(r[5] or 0),
                    int(r[6] or 0))

    @staticmethod
    def _row_to_consumable(r):
        return Consumable(str(r[0]), str(r[1]), str(r[2]),
                          int(r[3] or 0), int(r[4] or 10),
                          str(r[5] or "🧪") if len(r) > 5 else "🧪")

    @staticmethod
    def _row_to_event(r):
        import json as _json
        choices_str = str(r[3] or "[]") if len(r) > 3 else "[]"
        try:
            choices = _json.loads(choices_str)
        except Exception:
            choices = []
        return ExploreEvent(str(r[0]), str(r[1]), str(r[2]), choices)

    def _read_narratives(self, wb) -> list:
        """从 Excel「叙事」sheet 加载文字列表"""
        sheet_name = "叙事"
        if sheet_name not in wb.sheetnames:
            return []
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        result = []
        for row in rows:
            if row and row[0] is not None:
                text = str(row[0]).strip()
                if text:
                    result.append(text)
        return result


# 全局单例
_loader: Optional[DataLoader] = None


def get_loader() -> DataLoader:
    global _loader
    if _loader is None:
        _loader = DataLoader()
    return _loader
