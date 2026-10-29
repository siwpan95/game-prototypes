import os
import sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from pathlib import Path
from docx import Document

# 源 docx 与输出路径：命令行参数优先，其次环境变量，最后落在本仓 docs/ 下。
# （2026-09-23 去本地化：原硬编码 原路径已失效）
_HERE = Path(__file__).resolve().parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(os.environ.get("DOCX_SRC") or (_HERE / "docs" / "1.docx"))
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(os.environ.get("DOCX_OUT") or (_HERE / "docs" / "docx_full.txt"))
if not SRC.exists():
    sys.exit(f"[错误] 找不到源 docx：{SRC}\n  用法：python _export_docx.py [源.docx] [输出.txt]（或用 DOCX_SRC / DOCX_OUT 环境变量）")
OUT.parent.mkdir(parents=True, exist_ok=True)

doc = Document(str(SRC))
with open(OUT, 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            f.write(f'[{i}] {p.text}\n')
    for ti, table in enumerate(doc.tables):
        f.write(f'\n=== TABLE {ti} ===\n')
        for ri, row in enumerate(table.rows):
            cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            line = '  Row' + str(ri) + ': ' + ' | '.join(cells)
            f.write(line + '\n')
print('Done')
