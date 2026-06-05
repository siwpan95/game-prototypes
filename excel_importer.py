"""
excel_importer.py  ──  TwelveLabors Excel → JSON 导入工具
用法：
    python excel_importer.py              # 导入 TwelveLabors.xlsx（同目录）
    python excel_importer.py my.xlsx      # 导入指定文件
    python excel_importer.py --check      # 仅校验，不写文件
支持的 Sheet：
    characters      角色基础属性与起始遗物
    relics          遗物定义（含钩子）
    endings         结局条件
    random_events   随机事件（含选项）
    labours         试炼池（哪些章节属于哪个池）
    chapter_xxx     直接以章节ID命名的 Sheet，写入对应章节文件
"""

import json
import os
import sys
import re
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("[错误] 缺少 openpyxl，请先安装：pip install openpyxl")
    sys.exit(1)

BASE_DIR = Path(__file__).parent
DATA_DIR = BASE_DIR / "data"
CHAPTERS_DIR = DATA_DIR / "chapters"

# ────────────────────────────────────────────────────────────
# 通用工具
# ────────────────────────────────────────────────────────────

def _cell(row, col_map, key, default=None):
    """从行数据中取值，去除首尾空格，空字符串返回 default"""
    v = row.get(key, default)
    if v is None:
        return default
    v = str(v).strip() if not isinstance(v, str) else v.strip()
    if v == "" or v == "None":
        return default
    return v

def _int_cell(row, col_map, key, default=0):
    v = _cell(row, col_map, key)
    if v is None:
        return default
    try:
        return int(float(v))
    except (ValueError, TypeError):
        return default

def _json_cell(row, col_map, key, default=None):
    """解析 JSON 格式的单元格（遗物钩子、条件 JSON 等）"""
    v = _cell(row, col_map, key)
    if v is None:
        return default
    try:
        return json.loads(v)
    except json.JSONDecodeError as e:
        raise ValueError(f"列 [{key}] 的 JSON 解析失败：{v!r} → {e}")

def _tags_cell(row, col_map, key):
    """逗号分隔的 tag 列表"""
    v = _cell(row, col_map, key)
    if not v:
        return []
    return [t.strip() for t in v.split(",") if t.strip()]

def _effects_cell(row, col_map, key):
    """
    解析效果字符串，支持两种格式：
      1. JSON：{"荣耀":1,"生命":-1}
      2. 简写：荣耀+1,生命-2,神恩=0
    """
    v = _cell(row, col_map, key)
    if not v:
        return {}
    v = v.strip()
    if v.startswith("{"):
        return _json_cell(row, col_map, key, default={})
    result = {}
    for part in v.split(","):
        part = part.strip()
        if not part:
            continue
        m = re.match(r"^([\u4e00-\u9fa5A-Za-z_]+)\s*([+\-=])\s*(-?\d+)$", part)
        if m:
            stat, op, val = m.group(1), m.group(2), int(m.group(3))
            if op == "=":
                result[stat] = f"set_{val}" if val == 0 else val
            else:
                result[stat] = val if op == "+" else -abs(val)
        else:
            raise ValueError(f"无法解析效果格式：{part!r}，支持 JSON 或 简写如 '荣耀+1,生命-2'")
    return result

def sheet_to_rows(sheet):
    """将 Sheet 转为 [{列名: 值, ...}] 列表，跳过以 # 开头的注释行"""
    rows_data = list(sheet.iter_rows(values_only=True))
    if not rows_data:
        return []
    # 找到第一个非空非注释行作为表头
    header_idx = 0
    for i, row in enumerate(rows_data):
        first = str(row[0]).strip() if row[0] is not None else ""
        if first and not first.startswith("#"):
            header_idx = i
            break
    headers = [str(h).strip() if h is not None else f"_col{i}" for i, h in enumerate(rows_data[header_idx])]
    result = []
    for row in rows_data[header_idx + 1:]:
        first = str(row[0]).strip() if row[0] is not None else ""
        if not first or first.startswith("#"):
            continue
        result.append({headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers))})
    return result

# ────────────────────────────────────────────────────────────
# characters Sheet
# ────────────────────────────────────────────────────────────
# 列：id | name | desc | 荣耀 | 力量 | 智慧 | 神恩 | 疯狂 | 生命 | start_relic

def parse_characters(sheet):
    """返回 {char_id: {...}} dict"""
    rows = sheet_to_rows(sheet)
    result = {}
    stats = ["荣耀", "力量", "智慧", "神恩", "疯狂", "生命"]
    for i, row in enumerate(rows, 1):
        cid = _cell(row, None, "id")
        if not cid:
            raise ValueError(f"characters 第{i}行：id 不能为空")
        result[cid] = {
            "name": _cell(row, None, "name", cid),
            "desc": _cell(row, None, "desc", ""),
            "base_stats": {s: _int_cell(row, None, s, 0) for s in stats},
            "start_relic": _cell(row, None, "start_relic"),
        }
        if result[cid]["start_relic"] is None:
            del result[cid]["start_relic"]
    return result

# ────────────────────────────────────────────────────────────
# relics Sheet
# ────────────────────────────────────────────────────────────
# 列：id | name | rarity | tags | desc | on_acquire | hook_ON_STAT_CHECK | hook_ON_TAKE_DAMAGE | hook_ON_TURN_END | hook_ON_DEATH

RELIC_HOOK_COLS = {
    "hook_ON_STAT_CHECK":  "ON_STAT_CHECK",
    "hook_ON_TAKE_DAMAGE": "ON_TAKE_DAMAGE",
    "hook_ON_TURN_END":    "ON_TURN_END",
    "hook_ON_DEATH":       "ON_DEATH",
}

def parse_relics(sheet):
    rows = sheet_to_rows(sheet)
    result = {}
    for i, row in enumerate(rows, 1):
        rid = _cell(row, None, "id")
        if not rid:
            raise ValueError(f"relics 第{i}行：id 不能为空")
        relic = {
            "name": _cell(row, None, "name", rid),
            "rarity": _cell(row, None, "rarity", "common"),
            "tags": _tags_cell(row, None, "tags"),
        }
        desc = _cell(row, None, "desc")
        if desc:
            relic["desc"] = desc
        # on_acquire
        try:
            oa = _effects_cell(row, None, "on_acquire")
        except ValueError as e:
            raise ValueError(f"relics 第{i}行 [{rid}] on_acquire 解析错误：{e}")
        if oa:
            relic["on_acquire"] = oa
        # hooks
        hooks = {}
        for col, hook_key in RELIC_HOOK_COLS.items():
            v = _cell(row, None, col)
            if v:
                try:
                    hooks[hook_key] = json.loads(v)
                except json.JSONDecodeError as e:
                    raise ValueError(f"relics 第{i}行 [{rid}] {col} JSON错误：{e}\n内容：{v!r}")
        if hooks:
            relic["hooks"] = hooks
        result[rid] = relic
    return result

# ────────────────────────────────────────────────────────────
# endings Sheet
# ────────────────────────────────────────────────────────────
# 列：id | name | text | condition_json
# condition_json 示例：{"logic":"AND","rules":[{"type":"stat","name":"荣耀","op":">=","value":15}]}

def parse_endings(sheet):
    rows = sheet_to_rows(sheet)
    result = []
    for i, row in enumerate(rows, 1):
        eid = _cell(row, None, "id")
        if not eid:
            raise ValueError(f"endings 第{i}行：id 不能为空")
        try:
            cond = _json_cell(row, None, "condition_json")
        except ValueError as e:
            raise ValueError(f"endings 第{i}行 [{eid}] condition_json：{e}")
        result.append({
            "id": eid,
            "name": _cell(row, None, "name", eid),
            "condition": cond or {},
            "text": _cell(row, None, "text", ""),
        })
    return result

# ────────────────────────────────────────────────────────────
# random_events Sheet
# ────────────────────────────────────────────────────────────
# 表头（两行合并）：
#   event_id | title | text | conditions_json
#   option1_text | option1_effects | option1_req_stat | option1_req_val |
#                  option1_fail_text | option1_fail_effects | option1_branch | option1_relic
#   option2_xxx ... option3_xxx
# 最多支持 3 个选项列组

OPTION_FIELDS = ["text", "effects", "req_stat", "req_val", "fail_text", "fail_effects", "branch", "relic"]

def _build_option(row, prefix):
    """从行中提取一个选项，返回 dict 或 None（如果 text 为空）"""
    text = _cell(row, None, f"{prefix}_text")
    if not text:
        return None
    opt = {"text": text}
    try:
        eff = _effects_cell(row, None, f"{prefix}_effects")
    except ValueError as e:
        raise ValueError(f"选项 [{prefix}] effects 解析错误：{e}")
    opt["effects"] = eff

    req_stat = _cell(row, None, f"{prefix}_req_stat")
    req_val  = _int_cell(row, None, f"{prefix}_req_val", 0)
    if req_stat and req_val:
        opt["required_stat"] = {req_stat: req_val}

    fail_text = _cell(row, None, f"{prefix}_fail_text")
    if fail_text:
        opt["fail_text"] = fail_text
    try:
        fail_eff = _effects_cell(row, None, f"{prefix}_fail_effects")
    except ValueError as e:
        raise ValueError(f"选项 [{prefix}] fail_effects 解析错误：{e}")
    if fail_eff:
        opt["fail_effects"] = fail_eff

    branch = _cell(row, None, f"{prefix}_branch")
    if branch:
        opt["branch"] = branch

    relic = _cell(row, None, f"{prefix}_relic")
    if relic:
        opt["grant_relic"] = relic

    return opt

def parse_random_events(sheet):
    rows = sheet_to_rows(sheet)
    result = []
    for i, row in enumerate(rows, 1):
        eid = _cell(row, None, "event_id")
        if not eid:
            raise ValueError(f"random_events 第{i}行：event_id 不能为空")
        event = {
            "event_id": eid,
            "type": "random",
            "title": _cell(row, None, "title", ""),
            "text": _cell(row, None, "text", ""),
        }
        cond_raw = _cell(row, None, "conditions_json")
        if cond_raw:
            try:
                event["conditions"] = json.loads(cond_raw)
            except json.JSONDecodeError as e:
                raise ValueError(f"random_events 第{i}行 [{eid}] conditions_json 错误：{e}")

        options = []
        for n in ["1", "2", "3"]:
            try:
                opt = _build_option(row, f"option{n}")
            except ValueError as e:
                raise ValueError(f"random_events 第{i}行 [{eid}] option{n}：{e}")
            if opt:
                options.append(opt)
        if not options:
            raise ValueError(f"random_events 第{i}行 [{eid}]：至少需要一个选项（option1_text）")
        event["options"] = options
        result.append(event)
    return result

# ────────────────────────────────────────────────────────────
# labours Sheet
# ────────────────────────────────────────────────────────────
# 列：pool | chapter_id
# pool 值：physical / mental / divine / hidden / beowulf / kratos / ...

def parse_labours(sheet):
    rows = sheet_to_rows(sheet)
    result = {}
    for i, row in enumerate(rows, 1):
        pool = _cell(row, None, "pool")
        cid  = _cell(row, None, "chapter_id")
        if not pool or not cid:
            raise ValueError(f"labours 第{i}行：pool 和 chapter_id 不能为空")
        result.setdefault(pool, [])
        if cid not in result[pool]:
            result[pool].append(cid)
    return result

# ────────────────────────────────────────────────────────────
# chapter_xxx Sheet（直接写为章节事件列表）
# ────────────────────────────────────────────────────────────
# 列：event_id | type | title | text | (选项同 random_events)
# type 通常是 story

def parse_chapter_sheet(sheet):
    rows = sheet_to_rows(sheet)
    result = []
    for i, row in enumerate(rows, 1):
        eid = _cell(row, None, "event_id")
        if not eid:
            raise ValueError(f"第{i}行：event_id 不能为空")
        event = {
            "event_id": eid,
            "type": _cell(row, None, "type", "story"),
            "title": _cell(row, None, "title", ""),
            "text": _cell(row, None, "text", ""),
        }
        options = []
        for n in ["1", "2", "3"]:
            try:
                opt = _build_option(row, f"option{n}")
            except ValueError as e:
                raise ValueError(f"chapter 第{i}行 [{eid}] option{n}：{e}")
            if opt:
                # 检查 set_world_state
                ws_raw = _cell(row, None, f"option{n}_world_state")
                if ws_raw:
                    try:
                        opt["set_world_state"] = json.loads(ws_raw)
                    except json.JSONDecodeError as e:
                        raise ValueError(f"chapter 第{i}行 [{eid}] option{n}_world_state JSON错误：{e}")
                options.append(opt)
        if not options:
            raise ValueError(f"chapter 第{i}行 [{eid}]：至少需要一个选项")
        event["options"] = options
        result.append(event)
    return result

# ────────────────────────────────────────────────────────────
# 主导入函数
# ────────────────────────────────────────────────────────────

SHEET_HANDLERS = {
    "characters":    (parse_characters,    DATA_DIR / "characters.json",    "object"),
    "relics":        (parse_relics,         DATA_DIR / "relics.json",         "object"),
    "endings":       (parse_endings,        DATA_DIR / "endings.json",        "array"),
    "random_events": (parse_random_events,  DATA_DIR / "random_events.json",  "array"),
    "labours":       (parse_labours,        DATA_DIR / "labours.json",        "object"),
}

def import_excel(xlsx_path: str, check_only: bool = False) -> bool:
    """
    读取 xlsx_path，解析各 Sheet，写入对应 JSON 文件。
    check_only=True 时只校验不写文件。
    返回 True = 全部成功，False = 有错误。
    """
    path = Path(xlsx_path)
    if not path.exists():
        print(f"[错误] 文件不存在：{path}")
        return False

    print(f"{'[校验]' if check_only else '[导入]'} 读取 {path.name} ...")

    try:
        wb = openpyxl.load_workbook(str(path), data_only=True)
    except Exception as e:
        print(f"[错误] 无法打开 Excel 文件：{e}")
        return False

    has_error = False
    processed = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        sn_lower = sheet_name.lower().strip()

        # ── 已知固定 Sheet ──
        if sn_lower in SHEET_HANDLERS:
            handler, out_path, data_type = SHEET_HANDLERS[sn_lower]
            try:
                data = handler(sheet)
            except ValueError as e:
                print(f"  [错误] Sheet [{sheet_name}]：{e}")
                has_error = True
                continue

            if not check_only:
                out_path.parent.mkdir(parents=True, exist_ok=True)
                with open(out_path, "w", encoding="utf-8") as f:
                    json.dump(data, f, ensure_ascii=False, indent=2)
            count = len(data) if isinstance(data, (list, dict)) else "?"
            print(f"  [OK] {sheet_name} → {out_path.name}  ({count} 条{'，仅校验' if check_only else ''})")
            processed.append(sheet_name)

        # ── 章节 Sheet（以 chapter_ 开头）──
        elif sn_lower.startswith("chapter_"):
            chapter_id = sn_lower  # e.g. chapter_01
            out_path = CHAPTERS_DIR / f"{chapter_id}.json"
            try:
                data = parse_chapter_sheet(sheet)
            except ValueError as e:
                print(f"  [错误] Sheet [{sheet_name}]：{e}")
                has_error = True
                continue

            if not check_only:
                out_path.parent.mkdir(parents=True, exist_ok=True)
                with open(out_path, "w", encoding="utf-8") as f:
                    json.dump(data, f, ensure_ascii=False, indent=2)
            print(f"  [OK] {sheet_name} → chapters/{chapter_id}.json  ({len(data)} 事件{'，仅校验' if check_only else ''})")
            processed.append(sheet_name)

        # ── 忽略说明 Sheet ──
        elif sn_lower in ("readme", "说明", "guide", "help", "_readme"):
            print(f"  [跳过] {sheet_name}（说明页）")

        else:
            print(f"  [警告] 未识别的 Sheet 名称 [{sheet_name}]，已跳过（支持：characters/relics/endings/random_events/labours/chapter_xxx）")

    if not processed:
        print("[警告] 没有处理任何 Sheet，请检查 Sheet 名称是否正确。")
    elif not has_error:
        print(f"\n{'校验' if check_only else '导入'}完成，共处理 {len(processed)} 个 Sheet。✓")
    else:
        print(f"\n{'校验' if check_only else '导入'}完成，但存在错误，请修正后重试。✗")

    return not has_error


# ────────────────────────────────────────────────────────────
# CLI 入口
# ────────────────────────────────────────────────────────────

if __name__ == "__main__":
    args = sys.argv[1:]
    check_only = "--check" in args
    args = [a for a in args if not a.startswith("--")]

    if args:
        xlsx_file = args[0]
    else:
        # 默认查找同目录下的 TwelveLabors.xlsx
        xlsx_file = str(BASE_DIR / "TwelveLabors.xlsx")

    ok = import_excel(xlsx_file, check_only=check_only)
    sys.exit(0 if ok else 1)
