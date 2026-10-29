"""重新生成 config.xlsx 配置文件"""
import os, random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(BASE_DIR, "config.xlsx")

header_font = Font(bold=True, size=11, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="4472C4")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
cell_align = Alignment(vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

def style_data(ws, start_row, end_row, cols):
    for r in range(start_row, end_row + 1):
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = cell_align
            cell.border = thin_border

wb = Workbook()

# ===== Sheet1: 角色 =====
ws = wb.active
ws.title = "角色"
headers = ["属性", "值", "说明（可选）"]
rows = [
    ["姓名", "", "例如：内比布"],
    ["种族", "", "例如：半兽人，留空则随机"],
    ["职业", "", "例如：符文博学者，留空则随机"],
    ["等级", 1, ""],
    ["经验", 0, ""],
    ["力量", 10, ""],
    ["体质", 10, ""],
    ["敏捷", 10, ""],
    ["智力", 10, ""],
    ["智慧", 10, ""],
    ["魅力", 10, ""],
    ["HP", 100, ""],
    ["MP", 50, ""],
    ["章节", "", "例如：第一章 — 序章"],
    ["头衔", "", ""],
]
ws.append(headers)
for r in rows:
    ws.append(r)
style_header(ws, 1, 3)
style_data(ws, 2, len(rows) + 1, 3)
ws.column_dimensions["A"].width = 14
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 28

# ===== Sheet2: 装备 =====
from data.pq_equipment import generate_weapon, generate_shield, generate_armor
ws = wb.create_sheet("装备")
ws.append(["部位", "名称", "攻击", "防御"])
wn, wa = generate_weapon()
ws.append(["武器", wn, wa, 0])
sn, sd = generate_shield()
ws.append(["盾牌", sn, 0, sd])
for slot in ["头盔", "胸甲", "肩甲", "臂甲", "手套", "内甲", "裤甲", "腿甲", "足甲"]:
    an, ad = generate_armor()
    ws.append([slot, an, 0, ad])
for _ in range(3):
    ws.append(["", "", "", ""])
style_header(ws, 1, 4)
style_data(ws, 2, ws.max_row, 4)
ws.column_dimensions["A"].width = 10
ws.column_dimensions["B"].width = 35
ws.column_dimensions["C"].width = 10
ws.column_dimensions["D"].width = 10

# ===== Sheet3: 背包 =====
ws = wb.create_sheet("背包")
ws.append(["物品名", "数量"])
for name, qty in [("鼠尾巴", 3), ("骨片", 5), ("黏液", 2), ("尖牙", 4), ("麻袋", 1)]:
    ws.append([name, qty])
for _ in range(5):
    ws.append(["", ""])
style_header(ws, 1, 2)
style_data(ws, 2, ws.max_row, 2)
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 12

# ===== Sheet4: 任务 =====
ws = wb.create_sheet("任务")
ws.append(["名称", "目标类型", "目标", "所需数量", "奖励类型", "奖励值"])
quests = [
    ["收集鼠尾巴", "收集", "鼠尾巴", 5, "金币", "100"],
    ["击败狗头人", "击杀", "狗头人", 3, "装备", "随机装备"],
    ["收集骨片", "收集", "骨片", 8, "属性", "力量+2"],
    ["击败兽人", "击杀", "兽人", 2, "特殊宝物", "特殊宝物"],
    ["收集黏液", "收集", "黏液", 6, "法术", "随机法术"],
    ["击败巨魔", "击杀", "巨魔", 1, "金币", "250"],
    ["收集尖牙", "收集", "尖牙", 10, "属性", "体质+1"],
    ["击败食尸鬼", "击杀", "食尸鬼", 3, "装备", "随机装备"],
]
for q in quests:
    ws.append(q)
for _ in range(5):
    ws.append(["", "", "", "", "", ""])
style_header(ws, 1, 6)
style_data(ws, 2, ws.max_row, 6)
ws.column_dimensions["A"].width = 20
ws.column_dimensions["B"].width = 10
ws.column_dimensions["C"].width = 16
ws.column_dimensions["D"].width = 10
ws.column_dimensions["E"].width = 12
ws.column_dimensions["F"].width = 16

# ===== Sheet5: 怪物 =====
from data.pq_monsters import MONSTERS
ws = wb.create_sheet("怪物")
ws.append(["名称", "等级", "HP", "攻击", "防御", "经验", "掉落物（分号分隔）"])
rng = random.Random(42)
sample = rng.sample(MONSTERS, min(20, len(MONSTERS)))
for name, lv, loot in sorted(sample, key=lambda x: x[1]):
    ws.append([name, lv, int(30 + lv * 12), int(5 + lv * 2.5), int(2 + lv * 1.5), int(lv * 15), ";".join(loot)])
for _ in range(5):
    ws.append(["", "", "", "", "", "", ""])
style_header(ws, 1, 7)
style_data(ws, 2, ws.max_row, 7)
for col, w in zip("ABCDEFG", [22, 8, 8, 8, 8, 8, 40]):
    ws.column_dimensions[col].width = w

# ===== Sheet6: 章节 =====
from data.pq_acts import ACTS
ws = wb.create_sheet("章节")
ws.append(["名称", "最低等级", "怪物池（分号分隔）", "任务池（分号分隔）"])
for act in ACTS:
    ws.append([act["name"], act["min_level"], "", ""])
for _ in range(3):
    ws.append(["", "", "", ""])
style_header(ws, 1, 4)
style_data(ws, 2, ws.max_row, 4)
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 12
ws.column_dimensions["C"].width = 40
ws.column_dimensions["D"].width = 40

# ===== Sheet7: 法术 =====
from data.pq_spells import random_spell
ws = wb.create_sheet("法术")
ws.append(["名称", "类型", "威力", "法力消耗", "描述"])
for _ in range(8):
    sp = random_spell()
    ws.append([sp["name"], sp["spell_type"], sp["power"], sp["mp_cost"], sp["description"]])
for _ in range(5):
    ws.append(["", "", "", "", ""])
style_header(ws, 1, 5)
style_data(ws, 2, ws.max_row, 5)
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 10
ws.column_dimensions["C"].width = 8
ws.column_dimensions["D"].width = 10
ws.column_dimensions["E"].width = 30

# ===== Sheet8: 商店装备 =====
from data.pq_equipment import generate_shop_equipment
ws = wb.create_sheet("商店装备")
ws.append(["名称", "部位", "攻击", "防御", "价格"])
for _ in range(20):
    eq = generate_shop_equipment()
    ws.append([eq["name"], eq["slot"], eq["atk"], eq["def"], eq["price"]])
for _ in range(5):
    ws.append(["", "", "", "", ""])
style_header(ws, 1, 5)
style_data(ws, 2, ws.max_row, 5)
ws.column_dimensions["A"].width = 35
ws.column_dimensions["B"].width = 10
ws.column_dimensions["C"].width = 10
ws.column_dimensions["D"].width = 10
ws.column_dimensions["E"].width = 12

wb.save(OUTPUT)
print(f"config.xlsx 已生成 → {OUTPUT}")
print(f"  工作表: {wb.sheetnames}")
