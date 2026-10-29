"""
配置加载器 v2.1：从 config.xlsx 读取全部游戏配置。
支持旧版（5属性/旧装备槽）和新版（6属性/11槽/klass）混合读取。
新增：章节、法术、商店装备，装备带属性，任务带奖励。
所有 load_*() 共享同一份缓存，config.xlsx 只打开一次。
"""
import os
import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_FILE = os.path.join(BASE_DIR, "config.xlsx")

_config_cache = None

# 新版 11 装备槽
ALL_EQUIP_SLOTS = [
    "武器", "盾牌", "头盔", "胸甲", "肩甲",
    "臂甲", "手套", "内甲", "裤甲", "腿甲", "足甲",
]


def load_all_config() -> dict:
    """一次性打开 config.xlsx，读取全部 Sheet 数据，存入缓存后返回"""
    global _config_cache
    if _config_cache is not None:
        return _config_cache

    wb = openpyxl.load_workbook(CONFIG_FILE, data_only=True)

    # --- 角色 ---
    player = {}
    ws = wb["角色"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            key = str(row[0]).strip()
            player[key] = row[1]

    # 向后兼容：旧版只有 5 属性，新版有 6 属性（含 wisdom/智慧）
    # 如果 config 缺少智慧，不强制写入（交给 game_engine 的默认值处理）
    # 如果旧版用 "职业" 代替 "klass"，自动映射
    if "职业" in player and "klass" not in player:
        player["klass"] = player["职业"]

    # --- 装备（含攻击/防御属性，支持 11 槽）---
    equipment = {}
    ws = wb["装备"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        slot, name = row[0], row[1]
        if slot and name:
            slot = str(slot).strip()
            atk = int(row[2]) if len(row) > 2 and row[2] is not None else 0
            dfn = int(row[3]) if len(row) > 3 and row[3] is not None else 0
            equipment[slot] = {"name": str(name).strip(), "atk": atk, "def": dfn}

    # --- 背包 ---
    inventory = {}
    ws = wb["背包"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] and row[1] is not None:
            inventory[row[0]] = int(row[1])

    # --- 任务（含目标类型/奖励类型/奖励值）---
    quests = []
    ws = wb["任务"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        name, target_type, target, required = row[0], row[1], row[2], row[3]
        if not name or not target or required is None:
            continue
        reward_type = str(row[4]) if len(row) > 4 and row[4] else "金币"
        reward_value = str(row[5]) if len(row) > 5 and row[5] else "50"
        quests.append({
            "name": name,
            "target_type": str(target_type) if target_type else "收集",
            "target": target,
            "required": int(required),
            "reward_type": reward_type,
            "reward_value": reward_value,
        })

    # --- 怪物 ---
    monsters = []
    ws = wb["怪物"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        loot_str = str(row[6]) if row[6] else ""
        loot_table = [x.strip() for x in loot_str.split(";") if x.strip()]
        monsters.append({
            "name": row[0],
            "level": int(row[1]) if row[1] is not None else 1,
            "hp": int(row[2]) if row[2] is not None else 100,
            "attack": int(row[3]) if row[3] is not None else 10,
            "defense": int(row[4]) if row[4] is not None else 5,
            "xp": int(row[5]) if row[5] is not None else 50,
            "loot_table": loot_table,
        })

    # --- 章节 ---
    acts = []
    if "章节" in wb.sheetnames:
        ws = wb["章节"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            monster_pool = [x.strip() for x in str(row[2]).split(";") if x.strip()] if row[2] else []
            quest_pool = [x.strip() for x in str(row[3]).split(";") if x.strip()] if row[3] else []
            acts.append({
                "name": row[0],
                "min_level": int(row[1]) if row[1] else 1,
                "monster_pool": monster_pool,
                "quest_pool": quest_pool,
            })

    # --- 法术 ---
    spells = []
    if "法术" in wb.sheetnames:
        ws = wb["法术"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            spells.append({
                "name": row[0],
                "spell_type": str(row[1]) if row[1] else "伤害",
                "power": int(row[2]) if row[2] is not None else 20,
                "mp_cost": int(row[3]) if row[3] is not None else 10,
                "description": str(row[4]) if len(row) > 4 and row[4] else "",
            })

    # --- 商店装备 ---
    shop_items = []
    if "商店装备" in wb.sheetnames:
        ws = wb["商店装备"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            shop_items.append({
                "name": row[0],
                "slot": str(row[1]) if row[1] else "武器",
                "atk": int(row[2]) if row[2] is not None else 0,
                "def": int(row[3]) if row[3] is not None else 0,
                "price": int(row[4]) if row[4] is not None else 100,
            })

    wb.close()
    _config_cache = {
        "player": player,
        "equipment": equipment,
        "inventory": inventory,
        "quests": quests,
        "monsters": monsters,
        "acts": acts,
        "spells": spells,
        "shop_items": shop_items,
    }
    return _config_cache


# ==================== 兼容旧 API ====================

def load_player_config() -> dict:
    return load_all_config()["player"]


def load_equipment() -> dict:
    """返回 {部位: {name, atk, def}}"""
    return load_all_config()["equipment"]


def load_inventory() -> dict:
    return load_all_config()["inventory"]


def load_quests() -> list:
    """返回任务列表，每项含 reward_type / reward_value"""
    return load_all_config()["quests"]


def load_monsters() -> list:
    """返回怪物模板列表（纯数据 dict）"""
    return load_all_config()["monsters"]


def load_acts() -> list:
    """返回章节列表"""
    return load_all_config()["acts"]


def load_spells() -> list:
    """返回法术模板列表"""
    return load_all_config()["spells"]


def load_shop_items() -> list:
    """返回商店装备列表"""
    return load_all_config()["shop_items"]
