"""
Excel 实时输出模块 — 使用 openpyxl 直接写入，无 COM 依赖，彻底避免线程崩溃。
"""
import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font

from data.pq_equipment import EQUIPMENT_SLOTS

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_FILE = os.path.join(BASE_DIR, "progress_quest.xlsx")

ALL_SLOTS = [
    "武器", "盾牌", "头盔", "胸甲", "肩甲",
    "臂甲", "手套", "内甲", "裤甲", "腿甲", "足甲",
]


def write_game_to_excel(player, shop_catalog=None, filename=None):
    """将角色状态全部实时写入 Excel（openpyxl 纯文件模式）"""
    if filename is None:
        filename = DEFAULT_FILE

    # 尝试打开已有文件，否则新建
    if os.path.exists(filename):
        wb = openpyxl.load_workbook(filename)
    else:
        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

    # 确保所有 Sheet 存在
    required_sheets = [
        "角色面板", "装备", "背包", "任务", "法术书",
        "战斗日志", "城镇商店", "特殊宝物",
    ]
    for name in required_sheets:
        if name not in wb.sheetnames:
            wb.create_sheet(name)

    bold = Font(bold=True)
    gray = Font(color="646464")
    now_str = datetime.now().strftime("%H:%M:%S")

    # ==================== 角色面板 ====================
    ws = wb["角色面板"]
    ws.delete_rows(1, ws.max_row)

    _set(ws, "A1", "进度任务 v2 — 角色面板", bold, 14)
    title_str = f" {player.title}" if getattr(player, 'title', '') else ""
    race_val = getattr(player, 'race', '—')
    klass_val = getattr(player, 'klass', getattr(player, 'job', '—'))
    act_val = getattr(player, 'current_act', getattr(player, 'chapter', '第一章'))

    data = [
        ["姓名", player.name, "", "头衔", title_str.strip(), "", ""],
        ["种族", race_val, "", "职业", klass_val, "", ""],
        ["等级", str(player.level), "", "经验", f"{player.xp} / {player.xp_to_next}", "", ""],
        ["章节", str(act_val), "", "", "", "", ""],
        ["生命", f"{player.hp} / {player.max_hp}", "", "法力", f"{player.mp} / {player.max_mp}", "", ""],
        ["金币", str(player.gold), "", "", "", "", ""],
    ]
    for i, row in enumerate(data, 3):
        for j, val in enumerate(row):
            ws.cell(row=i, column=j + 1, value=val)

    r = 12
    _set(ws, f"A{r}", "— 属性 —", bold)
    r += 1
    attr_data = [
        ["力量", str(player.strength), "影响", "物理攻击力"],
        ["体质", str(player.constitution), "影响", "最大生命"],
        ["敏捷", str(player.dexterity), "影响", "闪避率 / 先手"],
        ["智力", str(player.intelligence), "影响", "最大法力 / 法术伤害"],
        ["智慧", str(player.wisdom), "影响", "法术防御 / 经验获取"],
        ["魅力", str(player.charisma), "影响", "商店折扣 / 特殊事件"],
    ]
    for i, row in enumerate(attr_data):
        for j, val in enumerate(row):
            ws.cell(row=r + i, column=j + 1, value=val)

    r += 7
    equip_atk = getattr(player, 'equip_atk', 0)
    equip_def = getattr(player, 'equip_def', 0)
    dodge = getattr(player, 'dodge_chance', 0)
    spell_def = getattr(player, 'spell_defense', 0)
    discount = getattr(player, 'shop_discount', 0)
    encumbrance = f"{getattr(player, 'inv_count', 0)} / {getattr(player, 'max_encumbrance', 50)}"

    _set(ws, f"A{r}", "— 战斗数值 —", bold)
    r += 1
    combat_data = [
        ["总攻击力", str(equip_atk), "", ""],
        ["总防御力", str(equip_def), "", ""],
        ["闪避率", f"{dodge:.1%}", "", ""],
        ["法术防御", str(spell_def), "", ""],
        ["商店折扣", f"{discount:.0%}", "", ""],
        ["负重", encumbrance, "", ""],
    ]
    for i, row in enumerate(combat_data):
        for j, val in enumerate(row):
            ws.cell(row=r + i, column=j + 1, value=val)

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 22

    # ==================== 装备 ====================
    ws = wb["装备"]
    ws.delete_rows(1, ws.max_row)
    _set(ws, "A1", "装备明细（11 槽）", bold, 12)
    for i, h in enumerate(["部位", "装备名称", "攻击", "防御"], 1):
        _set_hdr(ws, 3, i, h)
    for i, slot in enumerate(ALL_SLOTS, 4):
        item = player.equipment.get(slot)
        if item:
            ws.cell(row=i, column=1, value=slot)
            ws.cell(row=i, column=2, value=item.name)
            ws.cell(row=i, column=3, value=item.atk)
            ws.cell(row=i, column=4, value=item.def_)
        else:
            ws.cell(row=i, column=1, value=slot)
            ws.cell(row=i, column=2, value="（空）")
            ws.cell(row=i, column=3, value=0)
            ws.cell(row=i, column=4, value=0)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 10

    # ==================== 背包 ====================
    ws = wb["背包"]
    ws.delete_rows(1, ws.max_row)
    _set(ws, "A1", "背包 — 普通物品", bold, 12)
    _set_hdr(ws, 3, 1, "物品名称")
    _set_hdr(ws, 3, 2, "数量")
    inv_items = list(player.inventory.items())
    for i, (name, item) in enumerate(inv_items, 4):
        qty = item.quantity if hasattr(item, 'quantity') else item
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=qty)
    sr = 4 + len(inv_items) + 2
    _set(ws, f"A{sr}", "特殊宝物", bold, 12)
    _set_hdr(ws, sr + 1, 1, "宝物名称")
    _set_hdr(ws, sr + 1, 2, "价值（金币）")
    specials = getattr(player, 'specials', [])
    if specials:
        for i, s in enumerate(specials):
            ws.cell(row=sr + 2 + i, column=1, value=s.name)
            ws.cell(row=sr + 2 + i, column=2, value=s.value)
    else:
        ws.cell(row=sr + 2, column=1, value="（暂无特殊宝物）")
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 14

    # ==================== 任务 ====================
    ws = wb["任务"]
    ws.delete_rows(1, ws.max_row)
    for i, h in enumerate(["任务名称", "目标", "进度", "奖励", "状态"], 1):
        _set_hdr(ws, 1, i, h)
    for i, q in enumerate(player.quests, 2):
        ws.cell(row=i, column=1, value=q.name)
        ws.cell(row=i, column=2, value=q.target)
        ws.cell(row=i, column=3, value=f"'{getattr(q, 'current', 0)} / {q.required}")
        ws.cell(row=i, column=4, value=f"{getattr(q, 'reward_type', '?')}:{getattr(q, 'reward_value', '?')}")
        ws.cell(row=i, column=5, value="已完成" if getattr(q, 'completed', False) else "进行中")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 10

    # ==================== 法术书 ====================
    ws = wb["法术书"]
    ws.delete_rows(1, ws.max_row)
    for i, h in enumerate(["法术名称", "类型", "威力", "法力消耗", "描述"], 1):
        _set_hdr(ws, 1, i, h)
    for i, s in enumerate(player.spells, 2):
        ws.cell(row=i, column=1, value=s.name)
        ws.cell(row=i, column=2, value=s.spell_type)
        ws.cell(row=i, column=3, value=s.power)
        ws.cell(row=i, column=4, value=s.mp_cost)
        ws.cell(row=i, column=5, value=s.description)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 35

    # ==================== 战斗日志 ====================
    ws = wb["战斗日志"]
    ws.delete_rows(1, ws.max_row)
    _set_hdr(ws, 1, 1, "时间")
    _set_hdr(ws, 1, 2, "事件")
    for i, msg in enumerate(player.combat_log[-30:], 2):
        ws.cell(row=i, column=1, value=now_str)
        ws.cell(row=i, column=2, value=msg)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 65

    # ==================== 城镇商店 ====================
    ws = wb["城镇商店"]
    ws.delete_rows(1, ws.max_row)
    _set(ws, "A1", "城镇商店 — 装备目录", bold, 14)
    ws.cell(row=2, column=1, value=(
        f"回城触发：背包满（{getattr(player, 'inv_count', 0)}"
        f"/{getattr(player, 'max_encumbrance', 50)}）时自动回城 "
        "→ 卖出背包 → 购买更好装备"
    )).font = gray
    _set(ws, "A3",
         f"当前金币：{player.gold}   |   商店折扣：{getattr(player, 'shop_discount', 0):.0%}（魅力 {player.charisma}）",
         bold, 12)
    _set(ws, "A5", "商店目录：", bold)
    for i, h in enumerate(["部位", "装备名称", "攻击", "防御", "价格", "可购买"], 1):
        _set_hdr(ws, 6, i, h)
    if shop_catalog:
        sorted_shop = sorted(shop_catalog,
                             key=lambda x: (x.get("slot", ""), -(x.get("atk", 0) + x.get("def", 0))))
        for i, item in enumerate(sorted_shop, 7):
            ws.cell(row=i, column=1, value=item.get("slot", "?"))
            ws.cell(row=i, column=2, value=item.get("name", "?"))
            ws.cell(row=i, column=3, value=item.get("atk", 0))
            ws.cell(row=i, column=4, value=item.get("def", 0))
            ws.cell(row=i, column=5, value=item.get("price", 0))
            ws.cell(row=i, column=6, value="是" if item.get("price", 999) <= player.gold else "否")
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 8
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 10

    # ==================== 特殊宝物 ====================
    ws = wb["特殊宝物"]
    ws.delete_rows(1, ws.max_row)
    _set(ws, "A1", "特殊宝物 — 宝库", bold, 14)
    ws.cell(row=2, column=1, value="特殊宝物为任务奖励掉落，可在回城时出售换取金币。").font = gray
    for i, h in enumerate(["宝物名称", "基名", "前缀", "后缀", "价值（金币）"], 1):
        _set_hdr(ws, 4, i, h)
    if specials:
        for i, s in enumerate(specials, 5):
            ws.cell(row=i, column=1, value=s.name)
            ws.cell(row=i, column=2, value=getattr(s, 'base_name', '—'))
            ws.cell(row=i, column=3, value=getattr(s, 'attribute', '—'))
            ws.cell(row=i, column=4, value=getattr(s, 'of_name', '—'))
            ws.cell(row=i, column=5, value=s.value)
    else:
        ws.cell(row=5, column=1, value="（暂无特殊宝物 — 完成任务获取）").font = gray
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14

    # 保存
    wb.save(filename)


def _set(ws, cell, value, font=None, size=None):
    c = ws[cell]
    c.value = value
    if font:
        c.font = font
    if size:
        c.font = Font(size=size, bold=getattr(font, 'bold', False) if font else False)


def _set_hdr(ws, row, col, text):
    c = ws.cell(row=row, column=col, value=text)
    c.font = Font(bold=True)
