"""
generate_template.py  ──  生成 TwelveLabors.xlsx 模板（含现有游戏数据 + 填写说明）
用法：
    python generate_template.py
生成 TwelveLabors.xlsx 到项目根目录。
"""

import json
import os
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import (
        PatternFill, Font, Alignment, Border, Side
    )
    from openpyxl.utils import get_column_letter
except ImportError:
    print("[错误] 缺少 openpyxl，请先安装：pip install openpyxl")
    sys.exit(1)

BASE_DIR = Path(__file__).parent
DATA_DIR = BASE_DIR / "data"
OUT_PATH = BASE_DIR / "TwelveLabors.xlsx"

# ── 样式 ──────────────────────────────────────────────────────
COLOR_HEADER     = "1F3864"   # 深蓝
COLOR_HEADER_FG  = "FFFFFF"
COLOR_NOTE       = "FFF2CC"   # 淡黄（注释行）
COLOR_NOTE_FG    = "7F6000"
COLOR_EVEN       = "DEEAF1"   # 交替行浅蓝
COLOR_ODD        = "FFFFFF"
COLOR_BORDER     = "B8CCE4"

def _header_font():
    return Font(bold=True, color=COLOR_HEADER_FG, name="微软雅黑", size=10)

def _note_font():
    return Font(italic=True, color=COLOR_NOTE_FG, name="微软雅黑", size=9)

def _cell_font():
    return Font(name="微软雅黑", size=9)

def _header_fill():
    return PatternFill("solid", fgColor=COLOR_HEADER)

def _note_fill():
    return PatternFill("solid", fgColor=COLOR_NOTE)

def _even_fill():
    return PatternFill("solid", fgColor=COLOR_EVEN)

def _thin_border():
    side = Side(style="thin", color=COLOR_BORDER)
    return Border(left=side, right=side, top=side, bottom=side)

def _wrap():
    return Alignment(wrap_text=True, vertical="top")

def style_header(cell):
    cell.font   = _header_font()
    cell.fill   = _header_fill()
    cell.border = _thin_border()
    cell.alignment = _wrap()

def style_note(cell):
    cell.font   = _note_font()
    cell.fill   = _note_fill()
    cell.border = _thin_border()
    cell.alignment = _wrap()

def style_data(cell, row_idx):
    cell.font   = _cell_font()
    cell.fill   = _even_fill() if row_idx % 2 == 0 else PatternFill("solid", fgColor=COLOR_ODD)
    cell.border = _thin_border()
    cell.alignment = _wrap()

def write_sheet(ws, headers, notes, rows, col_widths=None):
    """
    headers : 列表，每个元素是列标题字符串
    notes   : 列表，每个元素是对应列的说明字符串（#注释行）
    rows    : 数据行列表（每行是列表）
    """
    # 注释行（第1行）
    ws.append(["# " + n for n in notes])
    for col_idx, _ in enumerate(headers, 1):
        style_note(ws.cell(1, col_idx))

    # 表头（第2行）
    ws.append(headers)
    for col_idx, _ in enumerate(headers, 1):
        style_header(ws.cell(2, col_idx))

    # 数据行
    for data_row in rows:
        ws.append(data_row)
        r = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            style_data(ws.cell(r, col_idx), r)

    # 列宽
    if col_widths:
        for col_idx, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = w
    ws.freeze_panes = "A3"   # 冻结前两行


def _load(rel):
    path = DATA_DIR / rel
    if path.exists():
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    return None

# ── Sheet: README ────────────────────────────────────────────
def make_readme(wb):
    ws = wb.create_sheet("说明")
    lines = [
        ("TwelveLabors.xlsx — 游戏内容配置说明",),
        ("",),
        ("本文件共包含以下 Sheet，每个 Sheet 对应一类游戏数据：",),
        ("",),
        ("  characters      ── 角色基础属性（荣耀/力量/智慧/神恩/疯狂/生命）与起始遗物",),
        ("  relics           ── 遗物定义，含稀有度、标签、获取效果、钩子（JSON格式）",),
        ("  endings          ── 结局条件（condition_json 为 JSON 格式，见示例）",),
        ("  random_events    ── 随机事件（最多3个选项，支持检定/失败效果/授予遗物）",),
        ("  labours          ── 试炼池（哪些章节属于physical/mental/divine/hidden/角色专属池）",),
        ("  chapter_xxx      ── 章节事件（Sheet名即为章节ID，如 chapter_01）",),
        ("",),
        ("使用方法：",),
        ("  1. 在对应 Sheet 中编辑或新增数据（勿修改表头）",),
        ("  2. 保存 Excel 文件",),
        ("  3. 在游戏目录运行：python excel_importer.py",),
        ("     仅校验不写文件：python excel_importer.py --check",),
        ("",),
        ("效果字段（effects/on_acquire/fail_effects）支持两种格式：",),
        ('  1. 简写：荣耀+1,生命-2,神恩=0   （=0 表示归零）',),
        ('  2. JSON：{"荣耀":1,"生命":-2}',),
        ("",),
        ("注释行：以 # 开头的行会被自动跳过，可用于写备注。",),
        ("",),
        ("condition_json 示例（endings Sheet）：",),
        ('  {"logic":"AND","rules":[{"type":"stat","name":"荣耀","op":">=","value":15},{"type":"counter","name":"labors_completed","op":"==","value":3}]}',),
        ('  type 可选：stat / counter / path_sequence / world_state',),
        ('  op   可选：>= / <= / == / > / <',),
        ("",),
        ("钩子类型（relics Sheet hook_xxx 列，填 JSON 对象）：",),
        ('  ON_STAT_CHECK   ── 属性检定时额外加成，示例：{"stat":"力量","bonus":2}',),
        ('  ON_TAKE_DAMAGE  ── 受伤时减免，示例：{"reduce_percent":0.3}',),
        ('  ON_TURN_END     ── 每回合结束时改变属性，示例：{"stat":"疯狂","change":1}',),
        ('  ON_DEATH        ── 死亡时触发，示例：{"action":"REVIVE_WITH_STAT","stat":"疯狂","value":5}',),
    ]
    for r, line in enumerate(lines, 1):
        ws.cell(r, 1, line[0])
        ws.cell(r, 1).font = Font(name="微软雅黑", size=10)
    ws.column_dimensions["A"].width = 100

# ── Sheet: characters ────────────────────────────────────────
def make_characters(wb):
    ws = wb.create_sheet("characters")
    headers = ["id","name","desc","荣耀","力量","智慧","神恩","疯狂","生命","start_relic"]
    notes   = [
        "角色ID（唯一，英文）",
        "显示名称（中文）",
        "简短描述",
        "初始荣耀值",
        "初始力量值",
        "初始智慧值",
        "初始神恩值",
        "初始疯狂值",
        "初始生命值",
        "起始遗物ID（relics Sheet 中的 id）"
    ]
    data = _load("characters.json") or {}
    rows = []
    for cid, c in data.items():
        s = c.get("base_stats", {})
        rows.append([
            cid,
            c.get("name",""),
            c.get("desc",""),
            s.get("荣耀",0), s.get("力量",0), s.get("智慧",0),
            s.get("神恩",0), s.get("疯狂",0), s.get("生命",0),
            c.get("start_relic",""),
        ])
    # 示例新角色行
    rows.append(["achilles","阿喀琉斯","特洛伊最强战士，命中注定的荣耀与悲剧。",8,13,4,3,2,7,""])
    widths = [14,12,28,6,6,6,6,6,6,16]
    write_sheet(ws, headers, notes, rows, widths)

# ── Sheet: relics ─────────────────────────────────────────────
def make_relics(wb):
    ws = wb.create_sheet("relics")
    headers = [
        "id","name","rarity","tags","desc","on_acquire",
        "hook_ON_STAT_CHECK","hook_ON_TAKE_DAMAGE","hook_ON_TURN_END","hook_ON_DEATH"
    ]
    notes = [
        "遗物ID（唯一，英文）",
        "显示名称",
        "稀有度：common/rare/epic/legendary/divine/cursed/chapter_unique",
        "标签，逗号分隔：防御/进攻/神圣/诅咒/智慧",
        "描述文本（可选）",
        "获取时效果：简写或JSON，如 力量+2,神恩-1",
        "检定钩子JSON：{\"stat\":\"力量\",\"bonus\":2}",
        "受伤钩子JSON：{\"reduce_percent\":0.3}",
        "回合结束钩子JSON：{\"stat\":\"疯狂\",\"change\":1}",
        "死亡钩子JSON：{\"action\":\"REVIVE_WITH_STAT\",\"stat\":\"疯狂\",\"value\":5}",
    ]
    data = _load("relics.json") or {}
    rows = []
    for rid, r in data.items():
        hooks = r.get("hooks", {})
        rows.append([
            rid,
            r.get("name",""),
            r.get("rarity","common"),
            ",".join(r.get("tags",[])),
            r.get("desc",""),
            json.dumps(r.get("on_acquire",{}), ensure_ascii=False) if r.get("on_acquire") else "",
            json.dumps(hooks.get("ON_STAT_CHECK",""), ensure_ascii=False) if "ON_STAT_CHECK" in hooks else "",
            json.dumps(hooks.get("ON_TAKE_DAMAGE",""), ensure_ascii=False) if "ON_TAKE_DAMAGE" in hooks else "",
            json.dumps(hooks.get("ON_TURN_END",""), ensure_ascii=False) if "ON_TURN_END" in hooks else "",
            json.dumps(hooks.get("ON_DEATH",""), ensure_ascii=False) if "ON_DEATH" in hooks else "",
        ])
    # 示例新遗物
    rows.append([
        "hermes_sandals","赫尔墨斯凉鞋","rare","神圣","风之神赐予的飞行凉鞋，足迹所至无所不能。",
        "荣耀+1,智慧+1","","","",""
    ])
    widths = [18,14,14,16,30,18,30,28,28,38]
    write_sheet(ws, headers, notes, rows, widths)

# ── Sheet: endings ────────────────────────────────────────────
def make_endings(wb):
    ws = wb.create_sheet("endings")
    headers = ["id","name","text","condition_json"]
    notes = [
        "结局ID（唯一，英文）",
        "结局名称（中文）",
        "结局描述文本",
        "触发条件（JSON），type可选 stat/counter/path_sequence/world_state",
    ]
    data = _load("endings.json") or []
    rows = []
    for e in data:
        rows.append([
            e.get("id",""),
            e.get("name",""),
            e.get("text",""),
            json.dumps(e.get("condition",{}), ensure_ascii=False),
        ])
    # 示例新结局
    rows.append([
        "ending_mad_god","疯神降临","你彻底失去理智，但获得了神明的力量。",
        '{"logic":"AND","rules":[{"type":"stat","name":"疯狂","op":">=","value":15}]}'
    ])
    widths = [20,14,40,60]
    write_sheet(ws, headers, notes, rows, widths)

# ── Sheet: random_events ──────────────────────────────────────
def make_random_events(wb):
    ws = wb.create_sheet("random_events")
    headers = [
        "event_id","title","text","conditions_json",
        "option1_text","option1_effects","option1_req_stat","option1_req_val",
        "option1_fail_text","option1_fail_effects","option1_branch","option1_relic",
        "option2_text","option2_effects","option2_req_stat","option2_req_val",
        "option2_fail_text","option2_fail_effects","option2_branch","option2_relic",
        "option3_text","option3_effects","option3_req_stat","option3_req_val",
        "option3_fail_text","option3_fail_effects","option3_branch","option3_relic",
    ]
    notes = [
        "事件ID（唯一）","标题","正文","触发条件JSON（可为空，表示无条件）",
        "选项1文本","选项1效果","选项1检定属性（可空）","选项1检定数值（可空）",
        "选项1失败文本（可空）","选项1失败效果（可空）","选项1分支标签（可空）","选项1授予遗物ID（可空）",
        "选项2文本","选项2效果","选项2检定属性","选项2检定数值",
        "选项2失败文本","选项2失败效果","选项2分支","选项2授予遗物",
        "选项3文本","选项3效果","选项3检定属性","选项3检定数值",
        "选项3失败文本","选项3失败效果","选项3分支","选项3授予遗物",
    ]
    data = _load("random_events.json") or []
    rows = []
    for ev in data:
        opts = ev.get("options", [])
        def get_opt(n):
            o = opts[n] if n < len(opts) else {}
            req = o.get("required_stat", {})
            req_stat = list(req.keys())[0] if req else ""
            req_val  = list(req.values())[0] if req else ""
            return [
                o.get("text",""),
                json.dumps(o.get("effects",{}), ensure_ascii=False) if o.get("effects") else "",
                req_stat, req_val,
                o.get("fail_text",""),
                json.dumps(o.get("fail_effects",{}), ensure_ascii=False) if o.get("fail_effects") else "",
                o.get("branch",""),
                o.get("grant_relic",""),
            ]
        cond_raw = json.dumps(ev["conditions"], ensure_ascii=False) if "conditions" in ev else ""
        rows.append([
            ev.get("event_id",""), ev.get("title",""), ev.get("text",""), cond_raw,
            *get_opt(0), *get_opt(1), *get_opt(2)
        ])
    # 示例新事件
    rows.append([
        "river_crossing","湍急的渡口","眼前的河流暴涨，两岸人群焦虑地等待着渡口重开。","",
        "跳入急流游过去","荣耀+2,生命-1","力量","8",
        "被急流冲走，撞上礁石。","生命-2","swim","",
        "寻找上游的浅滩绕行","智慧+1","","","","","detour","",
        "","","","","","","","",
    ])
    widths = [16,14,30,30] + [20,16,10,8,18,16,10,14]*3
    write_sheet(ws, headers, notes, rows, widths)

# ── Sheet: labours ────────────────────────────────────────────
def make_labours(wb):
    ws = wb.create_sheet("labours")
    headers = ["pool","chapter_id"]
    notes = [
        "试炼池名称：physical/mental/divine/hidden/beowulf/kratos/（自定义角色ID）",
        "章节ID（对应 chapters/ 目录下的文件名，不含 .json）",
    ]
    data = _load("labours.json") or {}
    rows = []
    for pool, chapters in data.items():
        for ch in chapters:
            rows.append([pool, ch])
    # 示例
    rows.append(["physical", "chapter_new"])
    widths = [16, 24]
    write_sheet(ws, headers, notes, rows, widths)

# ── Sheet: chapter_xxx ─────────────────────────────────────────
def make_chapter_sheet(wb, chapter_id, events):
    ws = wb.create_sheet(chapter_id)
    headers = [
        "event_id","type","title","text",
        "option1_text","option1_effects","option1_req_stat","option1_req_val",
        "option1_fail_text","option1_fail_effects","option1_branch","option1_relic","option1_world_state",
        "option2_text","option2_effects","option2_req_stat","option2_req_val",
        "option2_fail_text","option2_fail_effects","option2_branch","option2_relic","option2_world_state",
        "option3_text","option3_effects","option3_req_stat","option3_req_val",
        "option3_fail_text","option3_fail_effects","option3_branch","option3_relic","option3_world_state",
    ]
    notes = [
        "事件ID","类型(story)","标题","正文",
        "选1文本","选1效果","选1检定属性","选1检定值",
        "选1失败文本","选1失败效果","选1分支","选1遗物ID","选1世界状态JSON",
        "选2文本","选2效果","选2检定属性","选2检定值",
        "选2失败文本","选2失败效果","选2分支","选2遗物ID","选2世界状态JSON",
        "选3文本","选3效果","选3检定属性","选3检定值",
        "选3失败文本","选3失败效果","选3分支","选3遗物ID","选3世界状态JSON",
    ]
    rows = []
    for ev in events:
        opts = ev.get("options",[])
        def get_opt(n):
            o = opts[n] if n < len(opts) else {}
            req = o.get("required_stat",{})
            rs = list(req.keys())[0] if req else ""
            rv = list(req.values())[0] if req else ""
            ws_raw = json.dumps(o.get("set_world_state",{}), ensure_ascii=False) if o.get("set_world_state") else ""
            return [
                o.get("text",""),
                json.dumps(o.get("effects",{}), ensure_ascii=False) if o.get("effects") else "",
                rs, rv,
                o.get("fail_text",""),
                json.dumps(o.get("fail_effects",{}), ensure_ascii=False) if o.get("fail_effects") else "",
                o.get("branch",""),
                o.get("grant_relic",""),
                ws_raw,
            ]
        rows.append([
            ev.get("event_id",""), ev.get("type","story"),
            ev.get("title",""), ev.get("text",""),
            *get_opt(0), *get_opt(1), *get_opt(2)
        ])
    widths = [16,8,14,30] + [20,16,10,8,18,16,10,14,20]*3
    write_sheet(ws, headers, notes, rows, widths)


def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)   # 删掉默认 Sheet

    make_readme(wb)
    make_characters(wb)
    make_relics(wb)
    make_endings(wb)
    make_random_events(wb)
    make_labours(wb)

    # 导出现有章节（只导出 chapter_01/04/07/02/05/06/09/11/12 等已有章节）
    CHAPTERS_DIR = DATA_DIR / "chapters"
    if CHAPTERS_DIR.exists():
        for f in sorted(CHAPTERS_DIR.glob("chapter_*.json")):
            cid = f.stem
            try:
                with open(f, encoding="utf-8") as fp:
                    events = json.load(fp)
                make_chapter_sheet(wb, cid, events)
                print(f"  [Sheet] {cid}")
            except Exception as e:
                print(f"  [警告] 跳过 {cid}：{e}")

    wb.save(str(OUT_PATH))
    print(f"\n模板已生成：{OUT_PATH}")
    print("共包含以下 Sheet：", wb.sheetnames)


if __name__ == "__main__":
    main()
